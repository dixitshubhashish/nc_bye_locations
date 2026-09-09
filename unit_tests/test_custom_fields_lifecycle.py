"""Behavioral tests for the custom-field lifecycle and the gold view
definition-drift guard.

The rest of this batch's coverage is source-contract style (asserting on
inspect.getsource), matching the convention in test_reporting_cache.py and
test_mapping_layout_contract.py. These are the cases where a source assertion
is not enough - the behavior only shows up in which statements actually get
issued, so they drive the real functions against a fake BigQuery client.

Both the bigquery *module* and the *client* are faked. Faking only the client
looked sufficient (these tests passed standalone) but failed in a full-suite
run: other suites inject a fake google.cloud.bigquery into sys.modules and
leave the real namespace package degraded on cleanup, so a later
`from google.cloud import bigquery` resolved to a module with no
QueryJobConfig. See _fake_bigquery_modules() below.
"""

from __future__ import annotations

import sys
import types
import unittest
from types import SimpleNamespace
from unittest.mock import patch

import whitespace_tool.workflow_server as ws


def _fake_bigquery_modules():
    """Other suites here inject a fake google.cloud.bigquery into sys.modules
    and, on cleanup, leave the real namespace package degraded - a later
    `from google.cloud import bigquery` then resolves to a module missing
    QueryJobConfig. These tests passed standalone and failed only in a
    full-suite run because of it. Same fix as
    test_reporting_timeseries_cache.py's _FakeBigQueryModuleMixin: inject our
    own fake so ordering can't reach us."""
    fake_bigquery = types.SimpleNamespace(
        QueryJobConfig=lambda **kwargs: types.SimpleNamespace(**kwargs),
        ArrayQueryParameter=lambda name, type_, value: (name, type_, value),
        ScalarQueryParameter=lambda name, type_, value: (name, type_, value),
        SchemaField=lambda name, field_type, mode="NULLABLE", default_value_expression=None:
            types.SimpleNamespace(name=name, field_type=field_type, mode=mode,
                                  default_value_expression=default_value_expression),
        Table=lambda *a, **k: types.SimpleNamespace(schema=[]),
        LoadJobConfig=lambda **kwargs: types.SimpleNamespace(**kwargs),
    )
    fake_cloud = types.ModuleType("google.cloud")
    fake_cloud.bigquery = fake_bigquery
    return {
        "google": types.ModuleType("google"),
        "google.cloud": fake_cloud,
        "google.cloud.bigquery": fake_bigquery,
    }


class _FakeBigQueryModuleMixin(unittest.TestCase):
    def setUp(self) -> None:
        self._bq_mods = patch.dict(sys.modules, _fake_bigquery_modules())
        self._bq_mods.start()
        self.addCleanup(self._bq_mods.stop)


class FakeQueryJob:
    def __init__(self, rows=None, affected=0):
        self._rows = rows or []
        self.num_dml_affected_rows = affected

    def result(self):
        return list(self._rows)


class FakeClient:
    """Records every statement issued so a test can assert on what ran."""

    def __init__(self, query_results=None, affected_by_query=None):
        self.queries: list[str] = []
        self._query_results = query_results or {}
        self._affected = affected_by_query or {}
        self.tables: dict[str, object] = {}
        self.updated: list[tuple[str, list[str]]] = []
        self.loaded: list[list[dict]] = []

    def query(self, sql, job_config=None):
        self.queries.append(sql)
        for marker, rows in self._query_results.items():
            if marker in sql:
                return FakeQueryJob(rows)
        for marker, count in self._affected.items():
            if marker in sql:
                return FakeQueryJob(affected=count)
        return FakeQueryJob()

    def get_table(self, ref):
        if str(ref) in self.tables:
            return self.tables[str(ref)]
        raise _NotFound("missing")

    def update_table(self, table, fields):
        self.updated.append((getattr(table, "_ref", ""), list(fields)))
        return table

    def load_table_from_json(self, rows, _ref, job_config=None):
        self.loaded.append(list(rows))
        return FakeQueryJob()

    def create_table(self, table):
        return table


class _NotFound(Exception):
    code = 404


class CustomFieldArchiveTests(_FakeBigQueryModuleMixin):
    """Removing a custom field must archive it, never delete it.

    Listings already saved carry that field's values inside
    listings.custom_fields. Dropping the catalog row strands them with no
    label, type or provenance to interpret them by.
    """

    def _run_delete(self, client):
        with patch.object(ws, "_warehouse_settings", return_value=("p", "d", None)), \
             patch.object(ws, "_bigquery_client", return_value=client), \
             patch.object(ws, "invalidate_cache"):
            return ws.delete_custom_field({
                "field_key": "loyaltyTier", "password": "54321", "business_id": "b1",
            })

    def test_delete_issues_an_archive_update_and_never_a_delete(self) -> None:
        client = FakeClient(query_results={
            "SELECT field_id, label, is_custom": [
                {"field_id": "f1", "label": "Loyalty Tier", "is_custom": True, "business_id": "b1"},
            ],
        })
        result = self._run_delete(client)

        self.assertTrue(result["archived"])
        archive_statements = [q for q in client.queries if "is_archived = TRUE" in q]
        self.assertEqual(len(archive_statements), 1, "expected exactly one archive UPDATE")
        self.assertIn("UPDATE", archive_statements[0])
        self.assertIn("archived_at", archive_statements[0])
        # The whole point: nothing may be destroyed.
        self.assertFalse([q for q in client.queries if "DELETE FROM" in q.upper()],
                         "custom field removal must never issue a DELETE")

    def test_standard_fields_can_never_be_archived(self) -> None:
        client = FakeClient(query_results={
            "SELECT field_id, label, is_custom": [
                {"field_id": "f2", "label": "Address", "is_custom": False, "business_id": "b1"},
            ],
        })
        with self.assertRaisesRegex(ValueError, "Standard fields"):
            self._run_delete(client)
        self.assertFalse([q for q in client.queries if "is_archived = TRUE" in q])

    def test_missing_field_is_rejected_rather_than_silently_succeeding(self) -> None:
        client = FakeClient(query_results={"SELECT field_id, label, is_custom": []})
        with self.assertRaisesRegex(ValueError, "was not found"):
            self._run_delete(client)

    def test_admin_password_is_required(self) -> None:
        with self.assertRaisesRegex(ValueError, "Administrative password"):
            ws.delete_custom_field({"field_key": "x", "password": "wrong", "business_id": "b1"})


class CustomFieldCreateTests(_FakeBigQueryModuleMixin):
    """Re-creating a previously archived field restores that row instead of
    inserting a second row with the same slug - the archived definition
    already describes the values sitting in listings.custom_fields."""

    def _create(self, client, affected):
        client._affected = affected
        with patch.object(ws, "_warehouse_settings", return_value=("p", "d", None)), \
             patch.object(ws, "_bigquery_client", return_value=client), \
             patch.object(ws, "field_catalog", return_value=[]), \
             patch.object(ws, "_ensure_listings_table"), \
             patch.object(ws, "invalidate_cache"):
            return ws.create_custom_field({
                "label": "Loyalty Tier", "slug": "loyaltyTier",
                "password": "54321", "business_id": "b1", "type": "string",
            })

    def test_recreating_an_archived_field_revives_it_without_duplicating(self) -> None:
        client = FakeClient()
        result = self._create(client, {"SET is_archived = FALSE": 1})

        self.assertTrue(result["restored"])
        self.assertEqual(result["field"]["key"], "loyaltytier")
        # Revived, so no new catalog row may be written.
        self.assertEqual(client.loaded, [], "revive must not insert a duplicate slug")
        revive = [q for q in client.queries if "SET is_archived = FALSE" in q]
        self.assertEqual(len(revive), 1)
        self.assertIn("is_archived IS TRUE", revive[0])

    def test_a_genuinely_new_field_is_inserted(self) -> None:
        client = FakeClient()
        result = self._create(client, {"SET is_archived = FALSE": 0})

        self.assertNotIn("restored", result)
        self.assertEqual(len(client.loaded), 1, "a new field must be written to the catalog")
        self.assertEqual(client.loaded[0][0]["slug"], "loyaltytier")
        self.assertEqual(client.loaded[0][0]["table_name"], "listings")

    def test_creating_a_field_ensures_the_column_that_stores_its_values(self) -> None:
        # The catalog must not advertise somewhere for data to go that the
        # warehouse does not actually have.
        client = FakeClient()
        with patch.object(ws, "_warehouse_settings", return_value=("p", "d", None)), \
             patch.object(ws, "_bigquery_client", return_value=client), \
             patch.object(ws, "field_catalog", return_value=[]), \
             patch.object(ws, "invalidate_cache"), \
             patch.object(ws, "_ensure_listings_table") as ensure:
            client._affected = {"SET is_archived = FALSE": 0}
            ws.create_custom_field({
                "label": "Loyalty Tier", "slug": "loyaltyTier",
                "password": "54321", "business_id": "b1", "type": "string",
            })
        ensure.assert_called_once()

    def test_a_standard_field_name_cannot_be_hijacked(self) -> None:
        client = FakeClient()
        with patch.object(ws, "_warehouse_settings", return_value=("p", "d", None)), \
             patch.object(ws, "_bigquery_client", return_value=client):
            with self.assertRaisesRegex(ValueError, "standard field"):
                ws.create_custom_field({
                    "label": "Address", "slug": "address",
                    "password": "54321", "business_id": "b1", "type": "string",
                })


class GoldViewDefinitionDriftTests(_FakeBigQueryModuleMixin):
    """A gold view whose SELECT list changed in code still EXISTS with its old
    definition. _ensure_gold_reporting_views() used to return early on mere
    existence, so the edited view never reached an already-deployed
    environment and kept serving its old columns forever - live-confirmed
    when bronze and silver picked up custom_fields and gold did not.
    """

    REQUIRED = ("vw_zip_brand_activity", "vw_brand_summary", "vw_reporting_locations",
                "vw_reporting_filter_options", "vw_reporting_gap_base")

    def _client_with_views(self, description):
        client = FakeClient()
        for view in self.REQUIRED:
            client.tables[f"gold.{view}"] = SimpleNamespace(description=description, _ref=view)
        return client

    def test_views_stamped_with_the_current_version_are_left_alone(self) -> None:
        client = self._client_with_views(
            f"birdeye_gold_view_version={ws.GOLD_VIEW_DEFINITION_VERSION}")
        with patch.object(ws, "build_gold_layer", side_effect=AssertionError("must not rebuild")):
            self.assertFalse(ws._ensure_gold_reporting_views(client, "gold"))

    def test_a_stale_version_stamp_triggers_a_rebuild(self) -> None:
        client = self._client_with_views("birdeye_gold_view_version=1999-01-01.old")
        with patch.object(ws, "prepare_zipcodes"), \
             patch.object(ws, "build_silver_layer", return_value={"status": "ok"}), \
             patch.object(ws, "build_gold_layer", return_value={"views": []}) as rebuild, \
             patch.object(ws, "sync_gold_mirror"), \
             patch.object(ws, "_medallion_settings", return_value=("p", "b", "s", "g", None)), \
             patch.object(ws, "_ensure_dataset"), \
             patch.object(ws, "_ensure_businesses_table"), \
             patch.object(ws, "_ensure_listings_table"):
            ws._ensure_gold_reporting_views(client, "gold")
        rebuild.assert_called_once()

    def test_an_unstamped_legacy_view_is_treated_as_stale(self) -> None:
        # Views created before versioning existed carry no stamp at all.
        client = self._client_with_views(None)
        with patch.object(ws, "prepare_zipcodes"), \
             patch.object(ws, "build_silver_layer", return_value={"status": "ok"}), \
             patch.object(ws, "build_gold_layer", return_value={"views": []}) as rebuild, \
             patch.object(ws, "sync_gold_mirror"), \
             patch.object(ws, "_medallion_settings", return_value=("p", "b", "s", "g", None)), \
             patch.object(ws, "_ensure_dataset"), \
             patch.object(ws, "_ensure_businesses_table"), \
             patch.object(ws, "_ensure_listings_table"):
            ws._ensure_gold_reporting_views(client, "gold")
        rebuild.assert_called_once()

    def test_build_gold_layer_stamps_every_view_it_creates(self) -> None:
        # Without the stamp the drift check above can never pass, so every
        # ensure would rebuild the whole gold layer on every call.
        import inspect

        source = inspect.getsource(ws.build_gold_layer)
        self.assertIn('view_table.description = f"birdeye_gold_view_version={GOLD_VIEW_DEFINITION_VERSION}"', source)
        self.assertIn('client.update_table(view_table, ["description"])', source)
        self.assertIn("for view_ref in views:", source)


class ExtrasCollectionTests(unittest.TestCase):
    """_collect_extras() decides what survives the bronze write, so its edge
    cases are the difference between preserving a source column and losing
    it."""

    def test_mapped_columns_are_not_duplicated_into_extras(self) -> None:
        row = {"Name": "S", "Zip": "78701", "Loyalty": "gold"}
        fields = {"name": "Name", "postal_code": "Zip"}
        self.assertEqual(ws_normalization()._collect_extras(row, fields), {"Loyalty": "gold"})

    def test_internal_meta_key_is_never_exported(self) -> None:
        row = {"Loyalty": "gold", "__meta": {"template_id": "t1"}}
        self.assertEqual(ws_normalization()._collect_extras(row, {}), {"Loyalty": "gold"})

    def test_blank_and_null_values_are_skipped(self) -> None:
        row = {"A": "", "B": "   ", "C": None, "D": "keep"}
        self.assertEqual(ws_normalization()._collect_extras(row, {}), {"D": "keep"})

    def test_no_extras_returns_none_rather_than_an_empty_document(self) -> None:
        # None means the column stays NULL instead of storing a useless "{}".
        self.assertIsNone(ws_normalization()._collect_extras({"Name": "S"}, {"name": "Name"}))
        self.assertIsNone(ws_normalization()._collect_extras("not a dict", {}))

    def test_a_dotted_mapping_keeps_the_parent_object(self) -> None:
        # Deliberate trade: duplicating a little beats losing a column, since
        # never losing source data is the point of this field.
        row = {"address": {"street": "1 Main", "unit": "4B"}}
        self.assertEqual(ws_normalization()._collect_extras(row, {"address": "address.street"}),
                         {"address": {"street": "1 Main", "unit": "4B"}})


def ws_normalization():
    from whitespace_tool import normalization
    return normalization


class MetricExportWorkbookTests(unittest.TestCase):
    """The metrics sheet is computed from the sheet-1 rows, so the two sheets
    can never disagree and the headline figure is reproducible by hand."""

    def _load(self, metric, rows, reason=""):
        import io
        import openpyxl
        data = ws._metric_export_workbook(metric, rows, reason)
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_metrics_sheet_arithmetic_matches_the_entity_rows(self) -> None:
        rows = [
            {"listing_id": "L1", "brand": "A", "has_valid_zip": True},
            {"listing_id": "L2", "brand": "A", "has_valid_zip": False},
            {"listing_id": "L3", "brand": "B", "has_valid_zip": True},
            {"listing_id": "L4", "brand": "B", "has_valid_zip": True},
        ]
        wb = self._load("zip-completeness", rows)
        # Third sheet added: a per-brand competitor breakdown, since
        # whitespace analysis is a competitive question.
        self.assertEqual(wb.sheetnames, ["Listing Data", "Metrics", "Competitors"])
        self.assertEqual(wb["Listing Data"].max_row - 1, 4)

        metrics = {r[0]: (r[1], r[2]) for r in wb["Metrics"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(metrics["Rows in this export"][0], 4)
        self.assertEqual(metrics["Has Valid ZIP = TRUE"][0], 3)
        self.assertEqual(metrics["Has Valid ZIP = FALSE"][0], 1)
        self.assertIn("75.0%", metrics["Has Valid ZIP = TRUE"][1])
        self.assertEqual(metrics["Distinct Brand"][0], 2)

    def test_an_unavailable_source_is_labelled_not_presented_as_a_zero(self) -> None:
        wb = self._load("uncovered-zips", [], "view missing")
        metrics = [r[0] for r in wb["Metrics"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("DATA UNAVAILABLE", metrics)
        self.assertIn("view missing", str(wb["Listing Data"]["A2"].value))

    def test_a_genuine_zero_does_not_claim_unavailability(self) -> None:
        wb = self._load("uncovered-zips", [])
        metrics = [r[0] for r in wb["Metrics"].iter_rows(min_row=2, values_only=True)]
        self.assertNotIn("DATA UNAVAILABLE", metrics)

    def test_oversized_cells_are_truncated_rather_than_failing_the_export(self) -> None:
        # Excel refuses cells over 32,767 chars; raw_record JSON can exceed
        # that, and a hard failure would lose the whole export over one field.
        wb = self._load("invalid-listings", [{"raw_record": "x" * 40000}])
        value = wb["Listing Data"]["A2"].value
        self.assertLess(len(value), 32767)
        self.assertTrue(value.endswith("...[truncated]"))

    def test_every_row_key_becomes_a_column_even_when_rows_differ(self) -> None:
        # Different parses carry different columns; a union of keys means a
        # row's value is never dropped just because row 1 lacked that key.
        wb = self._load("invalid-listings", [{"a": 1}, {"b": 2}])
        headers = [c.value for c in wb["Listing Data"][1]]
        self.assertEqual(headers, ["A", "B"])


if __name__ == "__main__":
    unittest.main()


class MetricExportSqlShapeTests(unittest.TestCase):
    """Live-caught bug: the ZIP-family export computed `is_covered` in its
    SELECT list and then filtered on it in the WHERE of the SAME query, which
    BigQuery rejects with "Unrecognized name: is_covered". Every ZIP metric
    (Uncovered ZIPs, Covered Markets, Market ZIPs) returned a 400.

    No fake-client test could have caught this - a fake returns rows
    regardless of whether the SQL is valid. That is exactly why the live
    endpoint smoke run matters, and why this asserts on SQL shape.
    """

    def test_is_covered_is_computed_in_its_own_cte_before_being_filtered(self) -> None:
        import inspect
        import re

        source = inspect.getsource(ws.reporting_metric_export)
        self.assertIn("flagged AS (", source)
        self.assertIn("SELECT *, location_count > 0 AS is_covered FROM zips", source)
        self.assertIn("SELECT * FROM flagged", source)

        # The alias must never be defined and filtered inside one SELECT.
        zip_block = source.split("elif metric in _ZIP_METRIC_SLUGS:", 1)[1].split("elif metric in _BRAND_METRIC_SLUGS:", 1)[0]
        same_select = re.search(
            r"SELECT \*, location_count > 0 AS is_covered\s+FROM zips\s+\{coverage_filter\}", zip_block)
        self.assertIsNone(same_select, "is_covered filtered in the same SELECT that defines it")

    def test_every_zip_metric_has_a_coverage_filter_branch(self) -> None:
        import inspect

        source = inspect.getsource(ws.reporting_metric_export)
        self.assertIn('if metric == "uncovered-zips":', source)
        self.assertIn('coverage_filter = "WHERE NOT is_covered"', source)
        self.assertIn('coverage_filter = "WHERE is_covered"', source)
        # The universe metrics export both sides of the ratio.
        self.assertIn('coverage_filter = ""', source)


class MetricExportBundleTests(unittest.TestCase):
    """One click hands over a ZIP: the workbook, normalized CSVs, and a
    README naming the filters that produced it. Excel alone could carry
    neither the normalized copy nor the provenance note."""

    def _bundle(self, rows, reason="", filters=None):
        import io
        import zipfile
        data = ws._metric_export_bundle("zip-completeness", rows, reason, filters or {})
        return zipfile.ZipFile(io.BytesIO(data))

    def _rows(self):
        return [
            {"listing_id": "L1", "brand": "Domino's", "state_code": "TX", "city_name": "Austin",
             "zip_code": "78701", "has_valid_zip": True, "has_valid_coordinates": True},
            {"listing_id": "L2", "brand": "Domino's", "state_code": "CA", "city_name": "LA",
             "zip_code": "90001", "has_valid_zip": True, "has_valid_coordinates": False},
            {"listing_id": "L3", "brand": "Pizza Hut", "state_code": "TX", "city_name": "Dallas",
             "zip_code": "75201", "has_valid_zip": False, "has_valid_coordinates": True},
        ]

    def test_bundle_carries_workbook_normalized_csvs_and_a_readme(self) -> None:
        bundle = self._bundle(self._rows())
        self.assertEqual(
            sorted(bundle.namelist()),
            ["README.txt", "competitors.csv", "listings.csv", "metrics.csv", "zip-completeness.xlsx"])

    def test_readme_records_the_filters_that_produced_the_file(self) -> None:
        # A file with no provenance can't be trusted later - "is this all
        # states or just two?" has to be answerable from the file itself.
        bundle = self._bundle(self._rows(), filters={"state": ["TX", "CA"], "brand": []})
        readme = bundle.read("README.txt").decode()
        self.assertIn("state: TX, CA", readme)
        self.assertIn("Rows in this export: 3", readme)
        self.assertIn("Definition:", readme)

    def test_readme_says_so_when_nothing_was_filtered(self) -> None:
        readme = self._bundle(self._rows(), filters={}).read("README.txt").decode()
        self.assertIn("none (full population)", readme)

    def test_competitor_breakdown_reconciles_with_the_listing_rows(self) -> None:
        rows = self._rows()
        competitors = ws._metric_export_competitors(rows)
        self.assertEqual([c["brand"] for c in competitors], ["Domino's", "Pizza Hut"])
        self.assertEqual(sum(c["listings"] for c in competitors), len(rows))
        self.assertAlmostEqual(sum(c["share_pct"] for c in competitors), 100.0, places=1)
        dominos = competitors[0]
        self.assertEqual(dominos["listings"], 2)
        self.assertEqual(dominos["states"], 2)
        self.assertEqual(dominos["with_coordinates"], 1)

    def test_rows_without_a_brand_produce_no_competitor_sheet(self) -> None:
        # Better an explicit "nothing to compare" than a sheet of blanks.
        self.assertEqual(ws._metric_export_competitors([{"zip_code": "78701"}]), [])
        bundle = self._bundle([{"zip_code": "78701"}])
        self.assertNotIn("competitors.csv", bundle.namelist())

    def test_hitting_the_row_ceiling_is_reported_as_partial_not_unavailable(self) -> None:
        # A truncated file that looks complete is worse than no file.
        summary = ws._metric_export_summary(
            "zip-completeness", [{"a": 1}],
            f"This export reached the {ws.METRIC_EXPORT_ROW_LIMIT:,}-row ceiling, so it is a partial result.")
        labels = [row["metric"] for row in summary]
        self.assertIn("PARTIAL RESULT", labels)
        self.assertNotIn("DATA UNAVAILABLE", labels)

    def test_the_row_ceiling_covers_a_real_analyst_filter(self) -> None:
        # "20 states, up to 200k records" must actually come back.
        self.assertGreaterEqual(ws.METRIC_EXPORT_ROW_LIMIT, 200000)


class MetricExportFilterTests(unittest.TestCase):
    """Selecting 20 states used to collapse to a single scalar compared with
    `=`, so the export came back empty instead of covering the selection."""

    def test_filters_accept_repeated_and_comma_separated_values(self) -> None:
        import inspect
        source = inspect.getsource(ws.reporting_metric_export)
        assert 'def _multi(name: str, upper: bool = False)' in source
        assert 'for piece in str(raw or "").split(","):' in source
        assert 'bigquery.ArrayQueryParameter("states", "STRING", states)' in source
        assert 'bigquery.ArrayQueryParameter("brands", "STRING", brands)' in source

    def test_every_filter_uses_array_membership_not_equality(self) -> None:
        import inspect
        import re
        source = inspect.getsource(ws.reporting_metric_export)
        self.assertIn("IN UNNEST(@states)", source)
        self.assertIn("IN UNNEST(@brands)", source)
        # No scalar comparison may survive, or that filter silently matches
        # only the first selected value.
        self.assertEqual(re.findall(r"@(?:brand|state) = ''", source), [])


class EmptyArrayParameterTests(unittest.TestCase):
    """BigQuery converts an EMPTY array parameter to NULL, so
    `ARRAY_LENGTH(@brands) = 0` evaluates to NULL rather than TRUE - the whole
    WHERE clause becomes NULL and EVERY row is filtered out.

    Live-caught: an export with no filters at all returned 0 rows instead of
    37,208. Verified against real BigQuery:
        ARRAY_LENGTH(<empty param>) -> None,  (... = 0) -> None,  IS NULL -> True
    """

    def test_every_array_length_guard_in_the_export_is_null_safe(self) -> None:
        import inspect
        import re

        source = inspect.getsource(ws.reporting_metric_export)
        unguarded = re.findall(r"(?<!COALESCE\()ARRAY_LENGTH\(@(\w+)\) = 0", source)
        # Every occurrence must be wrapped in COALESCE(..., 0).
        bare = [name for name in unguarded
                if f"COALESCE(ARRAY_LENGTH(@{name}), 0) = 0" not in source]
        self.assertEqual(bare, [], f"unguarded empty-array checks: {bare}")
        self.assertIn("COALESCE(ARRAY_LENGTH(@brands), 0) = 0", source)
        self.assertIn("COALESCE(ARRAY_LENGTH(@states), 0) = 0", source)

    def test_no_bare_array_length_equals_zero_remains(self) -> None:
        import inspect
        import re

        source = inspect.getsource(ws.reporting_metric_export)
        # Strip the COALESCE-wrapped ones, then nothing may be left.
        stripped = source.replace("COALESCE(ARRAY_LENGTH(@brands), 0) = 0", "")
        stripped = stripped.replace("COALESCE(ARRAY_LENGTH(@states), 0) = 0", "")
        self.assertEqual(re.findall(r"ARRAY_LENGTH\(@\w+\) = 0", stripped), [])
